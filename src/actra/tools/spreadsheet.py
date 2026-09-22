from pathlib import Path
import openpyxl
from actra.models import ToolResult, SafetyLevel
from .registry import tool


@tool(
    name="create_workbook",
    description="Creates a new Excel (.xlsx) workbook with optional header columns.",
    parameters={
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "headers": {
                "type": "array",
                "items": {"type": "string"},
                "description": "List of header column names (e.g. ['Date', 'Item', 'Amount'])"
            },
            "sheet_name": {"type": "string", "description": "Title of initial worksheet (default 'Sheet1')"}
        },
        "required": ["path"]
    },
    safety=SafetyLevel.SAFE
)
def create_workbook(path: str, headers: list[str] | None = None, sheet_name: str = "Sheet1") -> ToolResult:
    try:
        file_path = Path(path).expanduser()
        file_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        if headers:
            ws.append(headers)
        wb.save(str(file_path))
        return ToolResult.ok("create_workbook", path=str(file_path), headers=headers or [])
    except Exception as e:
        return ToolResult.fail("create_workbook", "ERROR", str(e))


@tool(
    name="open_workbook",
    description="Opens and inspects an xlsx workbook, returning sheet names and dimensions.",
    parameters={
        "type": "object",
        "properties": {
            "path": {"type": "string"}
        },
        "required": ["path"]
    },
    safety=SafetyLevel.SAFE
)
def open_workbook(path: str) -> ToolResult:
    try:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            return ToolResult.fail("open_workbook", "NOT_FOUND", f"Workbook not found: {path}")
        
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        info = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            info[sheet_name] = {
                "max_row": ws.max_row,
                "max_column": ws.max_column
            }
        return ToolResult.ok("open_workbook", sheets=info)
    except Exception as e:
        return ToolResult.fail("open_workbook", "ERROR", str(e))


@tool(
    name="inspect_workbook",
    description="Detailed inspection of a workbook sheet, including headers and row count.",
    parameters={
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {"type": "string"}
        },
        "required": ["path"]
    },
    safety=SafetyLevel.SAFE
)
def inspect_workbook(path: str, sheet: str | None = None) -> ToolResult:
    try:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            return ToolResult.fail("inspect_workbook", "NOT_FOUND", f"Workbook not found: {path}")
        
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        
        headers = []
        if ws.max_row >= 1:
            for cell in ws[1]:
                headers.append(str(cell.value) if cell.value is not None else None)
        
        return ToolResult.ok(
            "inspect_workbook",
            sheet_name=ws.title,
            max_row=ws.max_row,
            max_column=ws.max_column,
            headers=headers
        )
    except Exception as e:
        return ToolResult.fail("inspect_workbook", "ERROR", str(e))


@tool(
    name="append_table_row",
    description="Appends a row to the table in the specified sheet based on a dictionary of values.",
    parameters={
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {"type": "string"},
            "values": {"type": "object"}
        },
        "required": ["path", "values"]
    },
    safety=SafetyLevel.SENSITIVE
)
def append_table_row(path: str, sheet: str | None = None, values: dict | None = None) -> ToolResult:
    values = values or {}
    try:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            return ToolResult.fail("append_table_row", "NOT_FOUND", f"Workbook not found: {path}")
        
        wb = openpyxl.load_workbook(str(file_path))
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        
        headers = []
        if ws.max_row >= 1:
            for cell in ws[1]:
                headers.append(str(cell.value) if cell.value is not None else None)
        else:
            return ToolResult.fail("append_table_row", "NO_HEADERS", "Sheet has no headers (row 1 is empty).")
            
        next_row = ws.max_row + 1
        for key, val in values.items():
            if key in headers:
                col_idx = headers.index(key) + 1
                ws.cell(row=next_row, column=col_idx, value=val)
        
        wb.save(str(file_path))
        return ToolResult.ok("append_table_row", row=next_row, values=values)
    except Exception as e:
        return ToolResult.fail("append_table_row", "ERROR", str(e))


@tool(
    name="read_sheet",
    description="Reads all rows and cells from an Excel sheet.",
    parameters={
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {"type": "string"}
        },
        "required": ["path"]
    },
    safety=SafetyLevel.SAFE
)
def read_sheet(path: str, sheet: str | None = None) -> ToolResult:
    try:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            return ToolResult.fail("read_sheet", "NOT_FOUND", f"Workbook not found: {path}")
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        return ToolResult.ok("read_sheet", sheet_name=ws.title, rows=rows)
    except Exception as e:
        return ToolResult.fail("read_sheet", "ERROR", str(e))


@tool(
    name="save_workbook",
    description="Saves a workbook to disk.",
    parameters={
        "type": "object",
        "properties": {
            "path": {"type": "string"}
        },
        "required": ["path"]
    },
    safety=SafetyLevel.SAFE
)
def save_workbook(path: str) -> ToolResult:
    try:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            return ToolResult.fail("save_workbook", "NOT_FOUND", f"Workbook not found: {path}")
        
        wb = openpyxl.load_workbook(str(file_path))
        wb.save(str(file_path))
        return ToolResult.ok("save_workbook", path=str(file_path))
    except Exception as e:
        return ToolResult.fail("save_workbook", "ERROR", str(e))


@tool(
    name="read_last_row",
    description="Reads the last row of a sheet for verification.",
    parameters={
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {"type": "string"}
        },
        "required": ["path"]
    },
    safety=SafetyLevel.SAFE
)
def read_last_row(path: str, sheet: str | None = None) -> ToolResult:
    try:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            return ToolResult.fail("read_last_row", "NOT_FOUND", f"Workbook not found: {path}")
        
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        
        last_row = ws.max_row
        row_data = []
        if last_row > 0:
            for cell in ws[last_row]:
                row_data.append(cell.value)
        
        return ToolResult.ok("read_last_row", row_index=last_row, data=row_data)
    except Exception as e:
        return ToolResult.fail("read_last_row", "ERROR", str(e))
